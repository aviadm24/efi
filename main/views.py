from django.shortcuts import render, redirect, render_to_response
from django.contrib import messages
from .forms import main_list_form, UploadFileForm, DateForm
from django.forms import formset_factory
from django.views.generic.list import ListView
from django.views.generic.edit import UpdateView, DeleteView, CreateView
from django.urls import reverse_lazy, reverse
from django.utils import timezone
from django.contrib.auth.models import User
from .models import main_list_model, Provider_data, Customer_data, From_data, To_data, Yeruka2_data, Yeruka_data,\
    Status_data, Service_data, Car_data, Driver_data, Flight_data, Fields_to_cancel
from django.http import JsonResponse
from itertools import chain
import json
import csv
import openpyxl
# from io import StringIO
import io
from django.db import IntegrityError
# import codecs
# https://stackoverflow.com/questions/20926403/heroku-rake-dbmigrate-results-in-error-r13-attach-error-failed-to-attach-t/21148716#21148716
# http://jsfiddle.net/QLfMU/116/
from django_tables2 import RequestConfig
from .tables import main_list_Table
from datetime import datetime, timedelta
from .admin import main_list_Resource
import os
from django.conf import settings
from django.core.files import File
from django.http import HttpResponse, StreamingHttpResponse
from .resources import main_list_Resource
from django.template.loader import render_to_string
from django.core.mail import EmailMessage
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required

BASE_DIR = settings.BASE_DIR


# def export_csv(request):
#     main_list_resource = main_list_Resource()
#     dataset = main_list_resource.export()
#     msg = EmailMessage("caneti", 'test csv attachment', to=['aviadm24@gmail.com'])
#     msg.attach('test.csv', dataset.csv, 'text/csv')
#     msg.content_subtype = "html"
#     msg.send()
#     return redirect('/')

def send_backup():
    main_list_resource = main_list_Resource()
    dataset = main_list_resource.export()

    msg = EmailMessage("test backup", 'test backup', to=['aviadm24@gmail.com'])
    msg.attach('test.csv', dataset.csv, 'text/csv')
    msg.content_subtype = "html"
    msg.send()


@csrf_exempt
def export_table(request):
    mainlist = request.POST.get('mainlist')
    # print('main list:', type(mainlist))
    sum_list = request.POST.get('sum_list')
    mail = request.POST.get('mail')
    content = render_to_string('main/send_tables.html', {'mainlist': mainlist, 'sum_list': sum_list}, request=request)
    msg = EmailMessage("caneti", content, to=[mail])
    # msg.attach('my_pdf.pdf', pdf, 'application/pdf')
    msg.content_subtype = "html"
    msg.send()

    return render(request, 'main/send_tables.html', {'mainlist': mainlist, 'sum_list': sum_list})


# def to_send(request):
#     print('to send view')
#     field_names = request.session.get('field_names', '')
#     # end_filter = request.session.get('end_filter', '')
#     return render(request, 'main/to_send.html', {'field_names': field_names}) #, 'end_filter': model.query
#
# def table_view(request):
#     table = main_list_Table(main_list_model.objects.all())
#     RequestConfig(request).configure(table)
#     return render(request, 'main/table_view.html', {'table': table})

@login_required(login_url='/login/')
def main_list(request):
    try:
        now = datetime.now()
        date = now.date()
        backup_date = request.session['back_up_date']
        print('try backup_date: ', datetime.strptime(backup_date, '%Y-%m-%d').date())
        delta = date - datetime.strptime(backup_date, '%Y-%m-%d').date()
        if delta.days > 30:
            print('30 days pased')
            send_backup()
            request.session['back_up_date'] = datetime.strftime(date, '%Y-%m-%d')
    except:
        now = datetime.now()
        date = now.date()
        request.session['back_up_date'] = datetime.strftime(date, '%Y-%m-%d')
        backup_date = request.session['back_up_date']
        print('except backup_date: ', backup_date)

    field_names = [f.name for f in main_list_model._meta.get_fields()]
    now = timezone.now()

    user = request.user
    print('user: ', user)
    print('user id: ', user.id)
    # https://stackoverflow.com/questions/2448978/how-to-limit-columns-returned-by-django-query
    user_main_list = main_list_model.objects.filter(user=request.user)
    # excluded_status_end = main_list_model.objects.exclude(Status='END')
    excluded_status_end = user_main_list.exclude(Status='END')
    excluded_status_end_and_past = excluded_status_end.exclude(Status='Past')

    #  https://stackoverflow.com/questions/7503241/django-models-selecting-single-field
    p_num_list = excluded_status_end_and_past.values_list('Project_num', flat=True)
    try:
        p_num_set = sorted(set(p_num_list), key=lambda k: int(k))
    except:
        p_num_set = set(p_num_list)
    print('sorted proj num: ', p_num_set)
    # p_num_filterd_list = main_list_model.objects.filter(Date__gte=now).values_list('Project_num', flat=True)
    customer_list = excluded_status_end_and_past.values_list('Customer', flat=True)
    customer_set = set(customer_list)
    # for i in customer_list:
    #     print('customer_list: ', i)
    provider_list = excluded_status_end_and_past.values_list('Provider', flat=True)
    provider_set = set(provider_list)
    last_month = datetime.today() - timedelta(days=30)
    today = datetime.today() - timedelta(days=1)

    # table_upcoming = main_list_Table(main_list_model.objects.filter(Date__gte=today))
    table_upcoming = main_list_Table(user_main_list.filter(Date__gte=today))
    # table_all = main_list_Table(main_list_model.objects.all())
    RequestConfig(request, paginate=False).configure(table_upcoming)

    if request.method == 'POST':
        print('main list - post')
        form = main_list_form(request.POST)
        date_form = DateForm(request.POST)
        # print('view - from:', request.POST['From'])
        if 'date_filter' in request.POST:
            pass
            # if date_form.is_valid():
            #     start = date_form.cleaned_data['start']
            #     end = date_form.cleaned_data['end']
            #     table_upcoming = main_list_Table(main_list_model.objects.filter(Date__range=[start, end]))
            #     return render(request, 'main/main_list_model_form.html', {'form': form,
            #                                                   'field_names': field_names[1:],
            #                                                   'table_upcoming': table_upcoming,
            #                                                   'p_num_list': p_num_set,
            #                                                   'customer_list': customer_set,
            #                                                   'provider_list': provider_set,
            #                                                   'date_form': date_form})
            # else:
            #     print('date form error: ', date_form.errors)
        else:
            if form.is_valid():
                for key, value in form.cleaned_data.items():
                    if key == 'Flight_num':
                        if Flight_data.objects.filter(Flight=value).exists():
                            pass
                        else:
                            f = Flight_data.objects.create(Flight=value)
                    if key == 'From':
                        if From_data.objects.filter(From=value).exists():
                            pass
                        else:
                            f = From_data.objects.create(From=value)
                    if key == 'To':
                        if To_data.objects.filter(To=value).exists():
                            pass
                        else:
                            f = To_data.objects.create(To=value)
                    if key == 'Driver_name':
                        driver_data, created = Driver_data.objects.get_or_create(Driver=value)
                        if created:
                            messages.success(request, 'New Driver_data was successfully created!')
                    if key == 'Provider':
                        provider_data, created = Provider_data.objects.get_or_create(Provider_name=value)
                        if created:
                            messages.success(request, 'New Provider_data was successfully created!')
                    if key == 'Contact' and value != '':
                        main_list_model.objects.filter(Project_num=form.cleaned_data['Project_num']).update(Contact=value)

                m_list_model = form.save(commit=False)
                # commit=False tells Django that "Don't send this to database yet.
                # I have more things I want to do with it."
                m_list_model.user = request.user  # Set the user object here
                m_list_model.save()  # Now you can send it to DB
                # form.save()
                messages.success(request, 'Your order was successfully updated!')
                return redirect('main_list')
            else:
                print('main list - got an error: ', form.errors)
                messages.error(request, 'Please correct the error below')

    else:
        form = main_list_form()
        date_form = DateForm()

    return render(request, 'main/main_list_model_form.html', {'form': form,
                                                              'field_names': field_names[1:],
                                                              'table_upcoming': table_upcoming,
                                                              'p_num_list': p_num_set,
                                                              'customer_list': customer_set,
                                                              'provider_list': provider_set,
                                                              'date_form': date_form})
@login_required(login_url='/login/')
def whole_list(request):
    user_main_list = main_list_model.objects.filter(user=request.user)
    excluded_status_end = user_main_list.exclude(Status__contains='END')
    excluded_status_end_and_past = excluded_status_end.exclude(Status__contains='Past')
    p_num_list = excluded_status_end_and_past.values_list('Project_num', flat=True)
    try:
        p_num_set = sorted(set(p_num_list), key=lambda k: int(k))
    except:
        p_num_set = set(p_num_list)
    print(p_num_set)
    customer_list = excluded_status_end_and_past.values_list('Customer', flat=True)
    customer_set = set(customer_list)
    provider_list = excluded_status_end_and_past.values_list('Provider', flat=True)
    provider_set = set(provider_list)
    # table_all = main_list_Table(main_list_model.objects.all())
    # table_all = main_list_Table(main_list_model.objects.exclude(Status='END'))
    table_all = main_list_Table(excluded_status_end_and_past)
    RequestConfig(request, paginate=False).configure(table_all)
    hidden_form = main_list_form()
    date_form = DateForm()
    return render(request, 'main/whole_list.html', {'hidden_form': hidden_form,
                                                    'table_all': table_all,
                                                    'p_num_list': p_num_set,
                                                    'customer_list': customer_set,
                                                    'provider_list': provider_set,
                                                    'date_form': date_form})

@login_required(login_url='/login/')
def staged_projects(request):
    print('stage_ended projects function')
    # staged_projects_filter = main_list_model.objects.filter(Client_status__contains='נשלחה חשבונית מס')
    user_main_list = main_list_model.objects.filter(user=request.user)
    staged_projects_filter = user_main_list.filter(Status__contains='Past')
    p_num_list = staged_projects_filter.values_list('Project_num', flat=True)
    print('end proj list: ', p_num_list)
    try:
        p_num_set = sorted(set(p_num_list), key=lambda k: int(k))
    except:
        p_num_set = set(p_num_list)
    customer_list = staged_projects_filter.values_list('Customer', flat=True)
    customer_set = set(customer_list)
    provider_list = staged_projects_filter.values_list('Provider', flat=True)
    provider_set = set(provider_list)
    # table_all = main_list_Table(main_list_model.objects.all())
    table_staged_projects = main_list_Table(staged_projects_filter)
    RequestConfig(request, paginate=False).configure(table_staged_projects)
    hidden_form = main_list_form()
    date_form = DateForm()
    return render(request, 'main/whole_list.html', {'hidden_form': hidden_form,
                                                    'table_all': table_staged_projects,
                                                    'p_num_list': p_num_set,
                                                    'customer_list': customer_set,
                                                    'provider_list': provider_set,
                                                    'date_form': date_form})

@login_required(login_url='/login/')
def ended_projects(request):
    print('ended projects function')
    user_main_list = main_list_model.objects.filter(user=request.user)
    ended_projects_filter = user_main_list.filter(Status__contains='END')
    p_num_list = ended_projects_filter.values_list('Project_num', flat=True)
    print('end proj list: ', p_num_list)
    try:
        p_num_set = sorted(set(p_num_list), key=lambda k: int(k))
    except:
        p_num_set = set(p_num_list)
    customer_list = ended_projects_filter.values_list('Customer', flat=True)
    customer_set = set(customer_list)
    provider_list = ended_projects_filter.values_list('Provider', flat=True)
    provider_set = set(provider_list)
    # table_all = main_list_Table(main_list_model.objects.all())
    table_ended_projects = main_list_Table(ended_projects_filter)
    RequestConfig(request, paginate=False).configure(table_ended_projects)
    hidden_form = main_list_form()
    date_form = DateForm()
    return render(request, 'main/whole_list.html', {'hidden_form': hidden_form,
                                                    'table_all': table_ended_projects,
                                                    'p_num_list': p_num_set,
                                                    'customer_list': customer_set,
                                                    'provider_list': provider_set,
                                                    'date_form': date_form})



def add_dollar(request):
    id = request.GET.get('id')
    td_id = request.GET.get('td_id')
    new_int = request.GET.get('new_int')
    data = main_list_model.objects.filter(pk=id)
    print('data: ', data.values())
    td_data = data.values()[0][td_id]
    print('new_int: ', new_int)

    if td_id == 'Extra_KM_client':
        main_list_model.objects.filter(pk=id).update(Extra_KM_client=new_int)
    if td_id == 'Extra_KM_provider':
        main_list_model.objects.filter(pk=id).update(Extra_KM_provider=new_int)
    if td_id == 'Cost_per_client':
        main_list_model.objects.filter(pk=id).update(Cost_per_client=new_int)
    if td_id == 'Cost_per_provider':
        main_list_model.objects.filter(pk=id).update(Cost_per_provider=new_int)
    if td_id == 'Cost_transfer_client':
        main_list_model.objects.filter(pk=id).update(Cost_transfer_client=new_int)

    if td_id == 'Cost_transfer_provider':
        main_list_model.objects.filter(pk=id).update(Cost_transfer_provider=new_int)
    if td_id == 'Cost_extra_hour_client':
        main_list_model.objects.filter(pk=id).update(Cost_extra_hour_client=new_int)
    if td_id == 'Cost_extra_hour_provider':
        main_list_model.objects.filter(pk=id).update(Cost_extra_hour_provider=new_int)
    if td_id == 'Cost_VIP_client':
        main_list_model.objects.filter(pk=id).update(Cost_VIP_client=new_int)
    if td_id == 'Cost_VIP_provider':
        main_list_model.objects.filter(pk=id).update(Cost_VIP_provider=new_int)
    if td_id == 'Cost_shonot_client':
        main_list_model.objects.filter(pk=id).update(Cost_shonot_client=new_int)
    if td_id == 'Cost_shonot_provider':
        main_list_model.objects.filter(pk=id).update(Cost_shonot_provider=new_int)

    return JsonResponse({})

def add_shekel(request):
    id = request.GET.get('id')
    td_id = request.GET.get('td_id')
    new_int = request.GET.get('new_int')
    data = main_list_model.objects.filter(pk=id)
    # print('data: ', data.values())
    td_data = data.values()[0][td_id]
    print('td_data: ', td_data)

    if td_id == 'Extra_KM_client':
        main_list_model.objects.filter(pk=id).update(Extra_KM_client=new_int)
    if td_id == 'Extra_KM_provider':
        main_list_model.objects.filter(pk=id).update(Extra_KM_provider=new_int)
    if td_id == 'Cost_per_client':
        main_list_model.objects.filter(pk=id).update(Cost_per_client=new_int)
    if td_id == 'Cost_per_provider':
        main_list_model.objects.filter(pk=id).update(Cost_per_provider=new_int)
    if td_id == 'Cost_transfer_client':
        main_list_model.objects.filter(pk=id).update(Cost_transfer_client=new_int)

    if td_id == 'Cost_transfer_provider':
        main_list_model.objects.filter(pk=id).update(Cost_transfer_provider=new_int)
    if td_id == 'Cost_extra_hour_client':
        main_list_model.objects.filter(pk=id).update(Cost_extra_hour_client=new_int)
    if td_id == 'Cost_extra_hour_provider':
        main_list_model.objects.filter(pk=id).update(Cost_extra_hour_provider=new_int)
    if td_id == 'Cost_VIP_client':
        main_list_model.objects.filter(pk=id).update(Cost_VIP_client=new_int)
    if td_id == 'Cost_VIP_provider':
        main_list_model.objects.filter(pk=id).update(Cost_VIP_provider=new_int)
    if td_id == 'Cost_shonot_client':
        main_list_model.objects.filter(pk=id).update(Cost_shonot_client=new_int)
    if td_id == 'Cost_shonot_provider':
        main_list_model.objects.filter(pk=id).update(Cost_shonot_provider=new_int)

    return JsonResponse({})

def add_euro(request):
    id = request.GET.get('id')
    td_id = request.GET.get('td_id')
    new_int = request.GET.get('new_int')
    data = main_list_model.objects.filter(pk=id)
    # print('data: ', data.values())
    td_data = data.values()[0][td_id]
    print('td_data: ', td_data)

    if td_id == 'Extra_KM_client':
        main_list_model.objects.filter(pk=id).update(Extra_KM_client=new_int)
    if td_id == 'Extra_KM_provider':
        main_list_model.objects.filter(pk=id).update(Extra_KM_provider=new_int)
    if td_id == 'Cost_per_client':
        main_list_model.objects.filter(pk=id).update(Cost_per_client=new_int)
    if td_id == 'Cost_per_provider':
        main_list_model.objects.filter(pk=id).update(Cost_per_provider=new_int)
    if td_id == 'Cost_transfer_client':
        main_list_model.objects.filter(pk=id).update(Cost_transfer_client=new_int)

    if td_id == 'Cost_transfer_provider':
        main_list_model.objects.filter(pk=id).update(Cost_transfer_provider=new_int)
    if td_id == 'Cost_extra_hour_client':
        main_list_model.objects.filter(pk=id).update(Cost_extra_hour_client=new_int)
    if td_id == 'Cost_extra_hour_provider':
        main_list_model.objects.filter(pk=id).update(Cost_extra_hour_provider=new_int)
    if td_id == 'Cost_VIP_client':
        main_list_model.objects.filter(pk=id).update(Cost_VIP_client=new_int)
    if td_id == 'Cost_VIP_provider':
        main_list_model.objects.filter(pk=id).update(Cost_VIP_provider=new_int)
    if td_id == 'Cost_shonot_client':
        main_list_model.objects.filter(pk=id).update(Cost_shonot_client=new_int)
    if td_id == 'Cost_shonot_provider':
        main_list_model.objects.filter(pk=id).update(Cost_shonot_provider=new_int)

    return JsonResponse({})


def add_color_json(request):
    color = request.GET.get('color')
    if color == '':
        id = request.GET.get('id')
        td_id = request.GET.get('td_id')
        text_color = request.GET.get('text_color')
        print('text_color: ', text_color)

        data = main_list_model.objects.filter(pk=id)
        color_data = data.values()[0]['Color']
        try:
            color_json_load = json.loads(color_data)
        except TypeError:
            color_json_load = json.loads('{}')
        print('color data: ', type(color_json_load))
        print('color: ', color_json_load)

        if text_color == 'true':
            color_json_load.pop(td_id + '_text', None)
            color_json = json.dumps(color_json_load)
            main_list_model.objects.filter(pk=id).update(Color=color_json)
            print('new text color: ', type(color_json))
        else:
            color_json_load.pop(td_id, None)
            color_json = json.dumps(color_json_load)
            main_list_model.objects.filter(pk=id).update(Color=color_json)
            print('new color: ', type(color_json))

    else:
        id = request.GET.get('id')
        td_id = request.GET.get('td_id')
        text_color = request.GET.get('text_color')
        print('text_color: ', text_color)

        data = main_list_model.objects.filter(pk=id)
        color_data = data.values()[0]['Color']
        try:
            color_json_load = json.loads(color_data)
        except TypeError:
            color_json_load = json.loads('{}')
        print('color data: ', type(color_json_load))
        if color_data == None:
            color_data = {}
        print('color: ', color_json_load)

        if text_color == 'true':
            color_json_load[td_id+'_text'] = color
            color_json = json.dumps(color_json_load)
            main_list_model.objects.filter(pk=id).update(Color=color_json)
            print('new text color: ', type(color_json))
        else:
            color_json_load[td_id] = color
            color_json = json.dumps(color_json_load)
            main_list_model.objects.filter(pk=id).update(Color=color_json)
            print('new color: ', type(color_json))

    return JsonResponse({'is_taken': 'is_taken'})


def stage_project_json(request):
    past_projects_json = request.GET.get('past_projects_json')
    past_projects = json.loads(past_projects_json)
    print('got past_projects_list: ', past_projects)
    main_list_model.objects.filter(Project_num__in=past_projects).update(Status='Past')
    return JsonResponse({})


def end_project_json(request):
    proj_num = request.GET.get('proj_num')
    print('changing proj num: ', proj_num)
    main_list_model.objects.filter(Project_num=proj_num).update(Status='END')
    return JsonResponse({})


def change_for_all_project_rows(request):
    class_name = request.GET.get('class_name')
    new_val = request.GET.get('new_val')
    proj_num = request.GET.get('proj_num')
    # very important
    # https: // stackoverflow.com / questions / 21797436 / django - how - to - update - model - field - from -json - data
    model_instances = main_list_model.objects.filter(Project_num=proj_num).exclude(Canceled='True')
    id_list = []
    for instance in model_instances:
        instance.update_field(class_name, new_val)
        instance.save()
        id_list.append(instance.id)
    return JsonResponse({'id_list': id_list})


# def cancel_currency_fields(request):
#     # https: // stackoverflow.com / questions / 31795295 / how - to - reference - a - set - of - model - fields - in -django
#     id = request.GET.get('id')
#     model_instance = main_list_model.objects.get(id=id)
#     print('model field canceled: ', model_instance.Canceled)
#     model_instance.Canceled = True
#     model_instance.save()
#     print('model field canceled: ', model_instance.Canceled)
#     # main_list_model.objects.filter(id=id).update(Canceled=True)
#     list_of_fields_to_cancel = Fields_to_cancel.objects.all().values_list('Currency_field', flat=True)
#     for Currency_field in list_of_fields_to_cancel:
#         model_instance.update_field(Currency_field, '0')
#         model_instance.save()
#     return JsonResponse({})




def update_cell_json(request):
    new_value = request.GET.get('new_value')
    print('new val', new_value)
    if new_value is not None:
        id = request.GET.get('id')
        td_id = request.GET.get('td_id')
        print('id val', id)
        td_id = td_id.split()[0]
        print('td_id val', td_id)
        model_instance = main_list_model.objects.get(pk=id)
        # td_id_data = data.values()[0][td_id]

        if td_id == 'Project_num':
            main_list_model.objects.filter(pk=id).update(Project_num=new_value)
        if td_id == 'Customer_num':
            main_list_model.objects.filter(pk=id).update(Customer_num=new_value)
        if td_id == 'Customer':
            main_list_model.objects.filter(pk=id).update(Customer=new_value)
        if td_id == 'Contact':
            main_list_model.objects.filter(pk=id).update(Contact=new_value)
            main_list_model.objects.filter(Project_num=getattr(model_instance, 'Project_num')).update(Contact=new_value)

        if td_id == 'Date':
            main_list_model.objects.filter(pk=id).update(Date=new_value)
        if td_id == 'Type_of_service':
            main_list_model.objects.filter(pk=id).update(Type_of_service=new_value)
        if td_id == 'Type_of_car':
            main_list_model.objects.filter(pk=id).update(Type_of_car=new_value)
        if td_id == 'Name':
            main_list_model.objects.filter(pk=id).update(Name=new_value)
        if td_id == 'Luggage':
            main_list_model.objects.filter(pk=id).update(Luggage=new_value)

        if td_id == 'Flight_num':
            main_list_model.objects.filter(pk=id).update(Flight_num=new_value)
            flight, created = Flight_data.objects.get_or_create(Flight=new_value)
            if created:
                messages.success(request, 'Your Flight_num was successfully created!')
        if td_id == 'Flight_shcedule':
            main_list_model.objects.filter(pk=id).update(Flight_shcedule=new_value)
        if td_id == 'Start_time':
            main_list_model.objects.filter(pk=id).update(Start_time=new_value)
        if td_id == 'End_time':
            main_list_model.objects.filter(pk=id).update(End_time=new_value)
        if td_id == 'From':
            main_list_model.objects.filter(pk=id).update(From=new_value)
            from_data, created = From_data.objects.get_or_create(From=new_value)
            if created:
                messages.success(request, 'Your From_data was successfully created!')
        if td_id == 'To':
            main_list_model.objects.filter(pk=id).update(To=new_value)
            to_data, created = To_data.objects.get_or_create(To=new_value)
            if created:
                messages.success(request, 'Your To_data was successfully created!')
        if td_id == 'Provider':
            main_list_model.objects.filter(pk=id).update(Provider=new_value)
            provider_data, created = Provider_data.objects.get_or_create(Provider_name=new_value)
            if created:
                messages.success(request, 'Your Provider_data was successfully created!')
        if td_id == 'Driver_name':
            main_list_model.objects.filter(pk=id).update(Driver_name=new_value)
            driver_data, created = Driver_data.objects.get_or_create(Driver=new_value)
            if created:
                messages.success(request, 'Your Driver_data was successfully created!')
        if td_id == 'Provider_status':
            main_list_model.objects.filter(pk=id).update(Provider_status=new_value)
        if td_id == 'Comments':
            print('comment new val: ', new_value)
            main_list_model.objects.filter(pk=id).update(Comments=new_value)
        if td_id == 'Status':
            # print('got status!!!!!!!!!!!!!!!!!!!')
            main_list_model.objects.filter(pk=id).update(Status=new_value)
            # print(main_list_model.objects.get(pk=id).values())
        if td_id == 'Client_status':
            main_list_model.objects.filter(pk=id).update(Client_status=new_value)

        if td_id == 'Extra_hours_client':
            main_list_model.objects.filter(pk=id).update(Extra_hours_client=new_value)
        if td_id == 'Based_on_client':
            main_list_model.objects.filter(pk=id).update(Based_on_client=new_value)
        if td_id == 'Extra_hours_provider':
            main_list_model.objects.filter(pk=id).update(Extra_hours_provider=new_value)
        if td_id == 'Based_on_provider':
            main_list_model.objects.filter(pk=id).update(Based_on_provider=new_value)

        if td_id == 'KM':
            main_list_model.objects.filter(pk=id).update(KM=new_value)
        if td_id == 'Extra_KM_client':
            main_list_model.objects.filter(pk=id).update(Extra_KM_client=new_value)
        if td_id == 'Extra_KM_provider':
            main_list_model.objects.filter(pk=id).update(Extra_KM_provider=new_value)
        if td_id == 'Cost_per_client':
            main_list_model.objects.filter(pk=id).update(Cost_per_client=new_value)

        if td_id == 'Cost_per_provider':
            main_list_model.objects.filter(pk=id).update(Cost_per_provider=new_value)
        if td_id == 'Cost_transfer_client':
            main_list_model.objects.filter(pk=id).update(Cost_transfer_client=new_value)
        if td_id == 'Cost_transfer_provider':
            main_list_model.objects.filter(pk=id).update(Cost_transfer_provider=new_value)
        if td_id == 'Cost_extra_hour_client':
            main_list_model.objects.filter(pk=id).update(Cost_extra_hour_client=new_value)

        if td_id == 'Cost_extra_hour_provider':
            main_list_model.objects.filter(pk=id).update(Cost_extra_hour_provider=new_value)
        if td_id == 'Cost_VIP_client':
            main_list_model.objects.filter(pk=id).update(Cost_VIP_client=new_value)
        if td_id == 'Cost_VIP_provider':
            main_list_model.objects.filter(pk=id).update(Cost_VIP_provider=new_value)
        if td_id == 'Cost_shonot_client':
            main_list_model.objects.filter(pk=id).update(Cost_shonot_client=new_value)
        if td_id == 'Cost_shonot_provider':
            main_list_model.objects.filter(pk=id).update(Cost_shonot_provider=new_value)
        if td_id == 'Canceled':
            main_list_model.objects.filter(pk=id).update(Canceled=new_value)
            if new_value == True:
                list_of_fields_to_cancel = Fields_to_cancel.objects.all().values_list('Currency_field', flat=True)
                for Currency_field in list_of_fields_to_cancel:
                    model_instance.update_field(Currency_field, '0')
                    model_instance.save()

    # else:
    #     id = request.GET.get('id')
    #     td_id = request.GET.get('td_id')
    #     text_color = request.GET.get('text_color')
    #     print('text_color: ', text_color)
    #
    #     data = main_list_model.objects.filter(pk=id)
    #     color_data = data.values()[0]['Color']
    #     try:
    #         color_json_load = json.loads(color_data)
    #     except TypeError:
    #         color_json_load = json.loads('{}')
    #     print('color data: ', type(color_json_load))
    #     if color_data == None:
    #         color_data = {}
    #     print('color: ', color_json_load)
    #
    #     if text_color == 'true':
    #         color_json_load[td_id+'_text'] = color
    #         color_json = json.dumps(color_json_load)
    #         main_list_model.objects.filter(pk=id).update(Color=color_json)
    #         print('new text color: ', type(color_json))
    #     else:
    #         color_json_load[td_id] = color
    #         color_json = json.dumps(color_json_load)
    #         main_list_model.objects.filter(pk=id).update(Color=color_json)
    #         print('new color: ', type(color_json))

    return JsonResponse({})


@login_required(login_url='/login/')
def upload_file(request):
    if request.method == 'POST':
        form = UploadFileForm(request.POST, request.FILES)
        if form.is_valid():
            # handle_uploaded_file(request.FILES['file'])
            csvfile = request.FILES['file']
            # dialect = csv.Sniffer().sniff(codecs.EncodedFile(csvfile, "utf-8").read(1024).decode('utf-8'))
            # https: // stackoverflow.com / questions / 29663749 / typeerror - cant - use - a - string - pattern - on - a - bytes - like - object - api
            # csvfile.open()
            # reader = csv.reader(codecs.EncodedFile(csvfile, "utf-8"), delimiter=',', dialect=dialect)

            # very important!!!
            # https://andromedayelton.com/2017/04/25/adventures-with-parsing-django-uploaded-csv-files-in-python3/
            csvfile.seek(0)
            # file_reader = csv.DictReader(io.StringIO(csvfile.read().decode('utf-8')))
            if csvfile.name == 'model_data.csv':
                file_reader = csv.reader(io.StringIO(csvfile.read().decode('utf-8')))
                for num, row in enumerate(file_reader):
                    # print('row: ',row)
                    if row[0] == '\ufeffType of service':
                        for data in row[1:]:
                            try:
                                Ser = Service_data()
                                Ser.Service = data
                                Ser.save()
                            except IntegrityError:
                                pass
                    elif row[0] == 'From':
                        for data in row[1:]:
                            try:
                                Fr = From_data()
                                Fr.From = data
                                Fr.save()
                            except IntegrityError:
                                pass
                    elif row[0] == 'To':
                        for data in row[1:]:
                            try:
                                T = To_data()
                                T.To = data
                                T.save()
                            except IntegrityError:
                                pass
                    elif row[0] == 'Yeruka':
                        for data in row[1:]:
                            try:
                                Yer = Yeruka_data()
                                Yer.Yeruka = data
                                Yer.save()
                            except IntegrityError:
                                pass
                    elif row[0] == 'Yeruka2':
                        for data in row[1:]:
                            try:
                                Yer2 = Yeruka2_data()
                                Yer2.Yeruka2 = data
                                Yer2.save()
                            except IntegrityError:
                                pass
                    elif row[0] == 'Status':
                        for data in row[1:]:
                            try:
                                Sta = Status_data()
                                Sta.Status = data
                                Sta.save()
                            except IntegrityError:
                                pass
                    elif row[0] == 'Type of car':
                        for data in row[1:]:
                            try:
                                Ca = Car_data()
                                Ca.Car = data
                                Ca.save()
                            except IntegrityError:
                                pass
            elif csvfile.name == 'csv_data.csv':
                file_reader = csv.DictReader(io.StringIO(csvfile.read().decode('utf-8')))
                for num, row in enumerate(file_reader):
                    if row['\ufeffשם ספק'] != '':
                        print('num: ', num, 'name: ',  row['\ufeffשם ספק'])
                        # try:
                        Provider = Provider_data()
                        print('name: ', row['\ufeffשם ספק'])
                        name = row['\ufeffשם ספק']
                        Provider.Provider_name = name
                        mail = row['כתובות מייל']
                        if mail == '':
                            mail = 'חסר'
                        Provider.email = mail
                        phone = row['טלפון']
                        if phone == '':
                            phone = 'חסר'
                        Provider.phone_num = phone
                        city = row['עיר']
                        if city == '':
                            city = 'חסר'
                        Provider.city = city
                        address = row['כתובת']
                        if address == '':
                            address = 'חסר'
                        Provider.address = address
                        id_num = row['מספר עוסק']
                        if id_num == '':
                            id_num = 'חסר'
                        print('id num: ', id_num)
                        Provider.id_num = id_num
                        contact = row['שם איש קשר']
                        if contact == '':
                            contact = 'חסר'
                        Provider.contact = contact
                        used_a_lot = row['ספק']
                        if used_a_lot == 'סס':
                            Provider.used_a_lot = True
                        Provider.save()
                        # except IntegrityError:
                        #     print('row: ', row)
                        #     print('IntegrityError: ', row['\ufeffשם ספק'])
            elif csvfile.name == 'csv_coxtumer_data.xlsx':
                print('file name: ', csvfile.name)
                file_reader = csv.DictReader(io.StringIO(csvfile.read().decode('utf-8')))
                for num, row in enumerate(file_reader):
                    print('row: ', num)
                    if row['\ufeffשם לקוח'] != '':
                        try:
                            Customer = Customer_data()
                            # print('name: ', row['\ufeffשם לקוח'])
                            name = row['\ufeffשם לקוח']
                            if '"' in name:
                                print('bad name: ', name)
                                name = '-'.join(name.split('"'))
                            Customer.Customer_name = name
                            mail = row['כתובות מייל']
                            if mail == '':
                                mail = 'חסר'
                            Customer.email = mail
                            phone = row['טלפון']
                            if phone == '':
                                phone = 'חסר'
                            Customer.phone_num = phone
                            city = row['עיר']
                            if city == '':
                                city = 'חסר'
                            Customer.city = city
                            address = row['כתובת']
                            if address == '':
                                address = 'חסר'
                            Customer.address = address
                            id_num = row['מספר עוסק']
                            if id_num == '':
                                id_num = 'חסר'
                            print('id num: ', id_num)
                            Customer.id_num = id_num
                            contact = row['שם איש קשר']
                            if contact == '':
                                contact = 'חסר'
                            Customer.contact = contact
                            used_a_lot = row['לקוח']
                            if used_a_lot == 'לל':
                                Customer.used_a_lot = True
                            Customer.save()
                        except IntegrityError:
                            print('IntegrityError: ', row['\ufeffשם לקוח'])
            else:
                print('file name: ', csvfile.name)
                # https: // www.pythoncircle.com / post / 591 / how - to - upload - and -process - the - excel - file - in -django /
                wb = openpyxl.load_workbook(csvfile, data_only=False)

                # getting a particular sheet by name out of many sheets
                worksheet = wb["Sheet1"]
                print(worksheet)
                excel_data = list()
                # iterating over the rows and
                # getting value from each cell in row

                # https: // stackoverflow.com / questions / 11464080 / django - model - field - by - variable

                for row_num, row in enumerate(worksheet.iter_rows()):
                    if row_num == 0:
                        # get user name
                        user = row[0].value
                        print('cell dict: ', dir(row[0]))
                        print('user: ', user)
                    elif row_num == 1:
                        # get and check field names
                        # https: // stackoverflow.com / questions / 3106295 / django - get - list - of - model - fields
                        model_field_names = [f.name for f in main_list_model._meta.get_fields()]
                        uploaded_model_field_names = [upf.value for upf in row]
                        # print(model_field_names)
                        # print(uploaded_model_field_names)
                        # print(model_field_names==uploaded_model_field_names)
                    else:
                        row_data = list()
                        row_id = row[0].value
                        model_instance, created = main_list_model.objects.get_or_create(pk=row_id)
                        for cell_num, cell in enumerate(row):
                            cell_val = cell.internal_value
                            # if type(cell_val) == 'float':
                            #     cell_val = int(cell_val)

                            if cell_num == 0:
                                pass
                            else:
                                model_field = model_field_names[cell_num]
                                if 'Cost' in model_field:
                                    if cell.has_style:
                                        cell_format = cell.number_format
                                        if type(cell_format) != str:
                                            print('type: ', type(cell_format))
                                            cell_format = str(cell_format)
                                        if '€' in cell_format:
                                            cell_val = '€'+str(cell_val)
                                            print('{} : {}'.format(cell_format, cell_val))
                                        elif '$' in cell_format:
                                            cell_val = '$'+str(cell_val)
                                            print('{} : {}'.format(cell_format, cell_val))
                                # https: // datatables.net / forums / discussion / 46914 / how - do - i - fix - the - dollar - symbol - appearing - at - the - end - of - amounts - in -exported - excel - spreadsheets
                                # https: // stackoverflow.com / questions / 34652300 / set - openpyxl - cell - format - to - currency
                                if cell_val == '—':
                                    cell_val = None
                                    if model_field == 'Comments':
                                        cell_val = ''
                                elif cell_val == 'NaN':
                                    cell_val = None
                                if model_field == 'Date':
                                    try:
                                        # cell_val = datetime.strptime(cell_val, '%m/%d/%Y').strftime('%Y-%m-%d')
                                        # 'l, F d, Y'
                                        # https: // ourcodeworld.com / articles / read / 555 / how - to - format - datetime - objects - in -the - view - and -template - in -django
                                        cell_val = datetime.strptime(cell_val, '%A, %B %d, %Y').strftime('%Y-%m-%d')

                                    except (ValueError, TypeError) as error:
                                        print('ValueError: ', cell_val)
                                        cell_val = None

                                if model_field == 'Flight_shcedule' or 'time' in model_field:
                                    try:
                                        cell_val = datetime.strptime(cell_val, '%m/%d/%Y %H:%M').strftime('%Y-%m-%d %H:%M')
                                    except (ValueError, TypeError) as error:
                                        print('ValueError: ', cell_val)
                                        cell_val = None
                                if cell_val == '✘':
                                    cell_val = False
                                elif cell_val == '✔':
                                    cell_val = True

                                # https: // stackoverflow.com / questions / 21797436 / django - how - to - update - model - field - from -json - data
                                # setattr(model_instance, model_field, cell_val)
                                if model_field == 'user':
                                    cell_val = User.objects.get(username=cell_val)
                                model_instance.update_field(model_field, cell_val)

                                # print('cell val: ', cell_val)
                            row_data.append(str(cell_val))
                        print('row id: ', row_id)
                        model_instance.save()
                        excel_data.append(row_data)

                return render(request, 'main/upload.html', {"excel_data": excel_data})

            # data = file_reader[0]

            # file_reader = csv.reader(csvfile, delimiter=',')
            # for row in file_reader:
            #     print(row)


            return render(request, 'main/upload.html', locals())
    else:
        form = UploadFileForm()
    return render(request, 'main/upload.html', {'form': form})

